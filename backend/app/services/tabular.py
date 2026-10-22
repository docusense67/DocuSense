"""Read CSV / Excel uploads into rows of strings for the import endpoints (Section 5B/5C)."""

import csv
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

# One import covers a campus-sized student list (Section 21: up to 5,000 students).
MAX_ROWS = 10_000
MAX_COLUMNS = 50
MAX_XLSX_EXPANDED_BYTES = 50_000_000


class TableFormatError(ValueError):
    pass


@dataclass
class RowErrors:
    """Per-row error report. `row` is the spreadsheet row number (header is row 1)."""

    rows: list[dict] = field(default_factory=list)

    def add(self, row: int, message: str) -> None:
        for entry in self.rows:
            if entry["row"] == row:
                entry["errors"].append(message)
                return
        self.rows.append({"row": row, "errors": [message]})

    def __bool__(self) -> bool:
        return bool(self.rows)


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(filename: str, content: bytes, required: set[str]) -> list[dict[str, str]]:
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise TableFormatError("CSV must be UTF-8 encoded") from None
        reader = csv.reader(io.StringIO(text))
        raw_rows = list(reader)
    elif name.endswith(".xlsx"):
        try:
            # .xlsx is a zip: a 2 MB upload can expand enormously, so check before parsing.
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded = sum(info.file_size for info in archive.infolist())
            if expanded > MAX_XLSX_EXPANDED_BYTES:
                raise TableFormatError("Excel file is too large once unpacked")
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            # max_col and the row cap bound memory even when the sheet declares a huge range.
            raw_rows = []
            for row in sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True):
                raw_rows.append([_cell_to_str(c) for c in row])
                if len(raw_rows) > MAX_ROWS + 1:
                    break
        except TableFormatError:
            raise
        except Exception:
            raise TableFormatError("Could not read the Excel file") from None
    else:
        raise TableFormatError("Upload a .csv or .xlsx file")

    while raw_rows and not any(cell.strip() for cell in raw_rows[-1]):
        raw_rows.pop()  # trailing blank rows (Excel pads to its declared range)
    if len(raw_rows) > MAX_ROWS + 1:
        raise TableFormatError(f"Too many rows: split the file into parts of {MAX_ROWS} rows")
    if not raw_rows:
        raise TableFormatError("File is empty")
    header = [h.strip().lower() for h in raw_rows[0]]
    missing = required - set(header)
    if missing:
        raise TableFormatError(f"Missing columns: {', '.join(sorted(missing))}")

    rows = []
    for values in raw_rows[1:]:
        cells = [_cell_to_str(v) for v in values]
        if not any(cells):
            rows.append({})  # keep numbering aligned with the spreadsheet
            continue
        rows.append({h: (cells[i] if i < len(cells) else "") for i, h in enumerate(header) if h})
    return rows
